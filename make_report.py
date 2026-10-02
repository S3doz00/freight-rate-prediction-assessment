from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


root = Path(__file__).resolve().parent
scorer_chart = root / "scorer_results" / "candidate_december.png"
chart = scorer_chart if scorer_chart.is_file() else root / "scorer_results" / "december_preview.png"
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(19)
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(8)
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(12.5)
styles["Heading 1"].font.color.rgb = RGBColor(6, 74, 86)
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(4)


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "064A56")
        tcPr.append(shd)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(8.5)
    return table


title = doc.add_paragraph("Freight Rate Prediction Assessment Report", style="Title")
title_border = OxmlElement("w:pBdr")
bottom_border = OxmlElement("w:bottom")
bottom_border.set(qn("w:val"), "nil")
title_border.append(bottom_border)
title._p.get_or_add_pPr().append(title_border)
doc.add_paragraph(
    "I trained on 48,000 labeled loads from January through October 2025 and produced rates for all "
    "12,000 unlabeled loads. Three forward monthly holdouts selected a robust regression model with "
    "a pooled mean absolute error of $104.37. The final labels are withheld, so this is a local "
    "validation result rather than the final submission score."
)

doc.add_paragraph("Data and quality checks", style="Heading 1")
doc.add_paragraph(
    "The training file has no duplicate load IDs or missing target rates. It contains 300 missing "
    "weights, 292 negative weights, and 374 missing market-index values. The final prediction file "
    "has 165 missing and 145 negative weights. Eight final-set cities are absent from training. "
    "I replaced missing and negative weights with the median positive training weight and added "
    "a missing-weight indicator. Origin and destination coordinates let the model generalize to "
    "new city names."
)
doc.add_paragraph(
    "Posted rates range from $57.22 to $25,533.00, with a 99th percentile of $5,972.83. "
    "A Huber fit reduces the influence of extreme rates while retaining every labeled row."
)

doc.add_paragraph("Training and validation approach", style="Heading 1")
doc.add_paragraph(
    "The final rows occur after the labeled period, so I used forward monthly holdouts. For the "
    "August holdout I trained through July; for September I trained through August; for October I "
    "trained through September. Feature cleaning was fit separately on each training period. "
    "I compared ordinary and Huber regression using common fields against versions that added "
    "market index, quote signal, and a date trend. I selected the model with the lowest pooled "
    "holdout mean absolute error, then refit it on all 48,000 labeled rows."
)

add_table(
    ["Holdout", "Rows", "Selected model MAE", "Selected model RMSE"],
    [
        ["August 2025", "4,759", "$99.49", "$614.10"],
        ["September 2025", "4,670", "$104.73", "$615.57"],
        ["October 2025", "4,853", "$108.82", "$648.36"],
        ["Pooled", "14,282", "$104.37", "$626.42"],
    ],
)
doc.add_paragraph(
    "The selected model uses distance, equipment, weight, origin and destination coordinates, "
    "day of week, and interactions. It omits market index, quote signal, and an extrapolated "
    "date trend because they did not improve the forward holdouts. The ordinary version of the "
    "same feature set had a pooled MAE of $110.15."
)

doc.add_page_break()
doc.add_paragraph("December fixed input prediction", style="Heading 1")
doc.add_paragraph(
    "The December chart holds the lane and load fixed: Lexington to Fort Wayne, 360 miles, "
    "Dry Van, and 32,000 lb. Only the date changes. The selected model can score these rows "
    "directly after filling the known city coordinates from training data. Its daily predictions "
    "range from $789.73 to $833.46."
)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(chart), width=Inches(6.75))
caption = doc.add_paragraph(
    "Figure 1. December 2025 daily predictions from the completed fixed-input CSV. "
    + ("Chart produced by the provided scorer." if chart == scorer_chart else
       "The provided scorer generates candidate_december.png from these same values.")
)
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in caption.runs:
    run.font.size = Pt(8.5)
    run.font.italic = True

doc.add_paragraph("Submission outputs", style="Heading 1")
doc.add_paragraph(
    "The model writes exactly one positive predicted rate for each validation load ID in "
    "validation_predictions.csv and fills all 31 daily December rows in "
    "data/december_chart_inputs.csv. The provided scorer checks both files and generates the "
    "required chart. It does not reveal the hidden final validation metrics."
)

doc.core_properties.title = "Freight Rate Prediction Assessment Report"
doc.core_properties.author = ""
doc.save(root / "report.docx")
